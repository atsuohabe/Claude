#!/usr/bin/env python3
"""
Excel (TOCFL 華語八千詞) → vocab JSON 変換スクリプト
使い方:
  python3 tools/convert_vocab.py
      --xlsx 華語八千詞表_master.xlsx   # Excelファイルのパス
      --sheet "準備級一級(Novice 1)"    # 対象シート名（省略時は全シート）
      --out  data/vocab-core.json      # 出力ファイル（省略時は自動分割）
"""
import argparse, json, re, sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl が必要です: pip install openpyxl")

# ────────────────────────────────────────────────
# TOCFL シート名 → difficulty 数値
# ────────────────────────────────────────────────
LEVEL_MAP: list[tuple[str, int]] = [
    ("準備級一級", 1), ("novice 1", 1), ("novice1", 1),
    ("準備級二級", 1), ("novice 2", 1), ("novice2", 1),
    ("入門級",     2), ("level 1",  2),
    ("基礎級",     2), ("level 2",  2),
    ("進階級",     3), ("level 3",  3),
    ("高階級",     4), ("level 4",  4),
    ("流利級",     5), ("level 5",  5),
]

# シート名 → tocfl_level（1〜7の一意な番号）
TOCFL_LEVEL_MAP: list[tuple[str, int]] = [
    ("準備級一級", 1), ("novice 1", 1), ("novice1", 1),
    ("準備級二級", 2), ("novice 2", 2), ("novice2", 2),
    ("入門級",     3), ("level 1",  3),
    ("基礎級",     4), ("level 2",  4),
    ("進階級",     5), ("level 3",  5),
    ("高階級",     6), ("level 4",  6),
    ("流利級",     7), ("level 5",  7),
]

def sheet_difficulty(sheet_name: str) -> int:
    key = sheet_name.lower()
    for pattern, level in LEVEL_MAP:
        if pattern.lower() in key:
            return level
    return 2

def sheet_tocfl_level(sheet_name: str) -> int:
    key = sheet_name.lower()
    for pattern, level in TOCFL_LEVEL_MAP:
        if pattern.lower() in key:
            return level
    return 0

# ────────────────────────────────────────────────
# 任務領域（コンテキスト）→ カテゴリ ID マッピング
# ────────────────────────────────────────────────
CONTEXT_TO_CATEGORY: dict[str, str] = {
    "個人資料":         "greetings",
    "與他人的關係":     "family",
    "房屋與家庭、環境": "home_furniture",
    "日常生活":         "daily_time",
    "飲食":             "food",
    "旅行":             "transport",
    "購物":             "shopping",
    "閒暇時間、娛樂":   "entertainment",
    "教育":             "education_work",
    "工作":             "education_work",
    "其他":             "greetings",
}

def context_to_category(context: str) -> str:
    return CONTEXT_TO_CATEGORY.get(context.strip(), "greetings")

# 注音（ボポモフォ）括弧注記を除去: 例 "名字(˙ㄗ)" → "名字"
_BOPOMOFO_RE = re.compile(r'[（(][˙ˊˇˋ\u02C9ㄅ-ㄩ\uF8F0-\uF8FF]+[）)]')

def strip_bopomofo(text: str) -> str:
    return _BOPOMOFO_RE.sub('', text).strip()

# ────────────────────────────────────────────────
# 詞類（品詞）→ 英語 POS マッピング
# ────────────────────────────────────────────────
POS_MAP: dict[str, str] = {
    "n":    "noun",
    "v":    "verb",
    "adj":  "adjective",
    "adv":  "adverb",
    "prep": "preposition",
    "conj": "conjunction",
    "det":  "determiner",
    "intj": "interjection",
    "m":    "measure_word",
    "num":  "number",
    "pron": "pronoun",
    "part": "particle",
    "aux":  "auxiliary",
}

def normalize_pos(raw: str) -> str:
    return POS_MAP.get(raw.strip().lower(), "noun")

# ────────────────────────────────────────────────
# 声調記号 → 数字声調
# ────────────────────────────────────────────────
DIACRITIC_TABLE: dict[str, tuple[str, str]] = {
    'ā':('a','1'), 'á':('a','2'), 'ǎ':('a','3'), 'à':('a','4'),
    'ē':('e','1'), 'é':('e','2'), 'ě':('e','3'), 'è':('e','4'),
    'ī':('i','1'), 'í':('i','2'), 'ǐ':('i','3'), 'ì':('i','4'),
    'ō':('o','1'), 'ó':('o','2'), 'ǒ':('o','3'), 'ò':('o','4'),
    'ū':('u','1'), 'ú':('u','2'), 'ǔ':('u','3'), 'ù':('u','4'),
    'ǖ':('ü','1'), 'ǘ':('ü','2'), 'ǚ':('ü','3'), 'ǜ':('ü','4'),
}

def to_pinyin_tones(pinyin: str) -> str:
    """nǐ hǎo → ni3 hao3"""
    result = []
    for syllable in pinyin.split():
        tone = "5"
        cleaned = ""
        for ch in syllable:
            if ch in DIACRITIC_TABLE:
                base, t = DIACRITIC_TABLE[ch]
                cleaned += base
                tone = t
            else:
                cleaned += ch
        result.append(cleaned + tone)
    return " ".join(result)

# ────────────────────────────────────────────────
# ヘッダー列検出（改行入りヘッダー対応）
# ────────────────────────────────────────────────
# ヘッダーセルが "任務領域\nContext" のように改行を含む場合に対応するため
# 各行の全トークン（\n・空白で分割）と比較する

def _header_tokens(cell_str: str) -> list[str]:
    """'任務領域\nContext' → ['任務領域', 'context']"""
    tokens = []
    for part in cell_str.replace('\n', ' ').split():
        tokens.append(part.lower().strip())
    return tokens

def find_col(header: list[str], aliases: list[str]) -> int | None:
    for alias in aliases:
        al = alias.lower().strip()
        for i, h in enumerate(header):
            tokens = _header_tokens(h)
            if al in tokens or al == h.lower().strip():
                return i
    return None

# ────────────────────────────────────────────────
# 列エイリアス定義
# ────────────────────────────────────────────────
CONTEXT_ALIASES = ["任務領域", "context", "カテゴリ", "category", "分類"]
HANZI_ALIASES   = ["詞彙", "vocabulary", "漢字", "繁体字", "中文", "華語", "詞語", "hanzi"]
PINYIN_ALIASES  = ["漢語拼音", "pinyin", "拼音", "ピンイン"]
POS_ALIASES     = ["詞類", "parts of speech", "part of speech", "pos", "品詞"]
JA_ALIASES      = ["日本語", "japanese", "meaning_ja", "日本語意味", "意味(日)", "和訳"]
EN_ALIASES      = ["英語", "english", "meaning_en", "英語意味", "英訳"]

# ────────────────────────────────────────────────
# シート解析
# ────────────────────────────────────────────────

def parse_sheet(sheet, difficulty: int, start_id: int, tocfl_level: int = 0) -> list[dict]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    # ヘッダー行を探す（最初の5行以内）
    header_idx = 0
    for idx, row in enumerate(rows[:5]):
        tokens_in_row = []
        for c in row:
            if c:
                tokens_in_row.extend(_header_tokens(str(c)))
        if any(a.lower() in tokens_in_row for a in ["詞彙", "vocabulary", "漢字", "hanzi"]):
            header_idx = idx
            break

    header = [str(c).strip() if c else "" for c in rows[header_idx]]

    ci_context = find_col(header, CONTEXT_ALIASES)
    ci_hanzi   = find_col(header, HANZI_ALIASES)
    ci_pinyin  = find_col(header, PINYIN_ALIASES)
    ci_pos     = find_col(header, POS_ALIASES)
    ci_ja      = find_col(header, JA_ALIASES)
    ci_en      = find_col(header, EN_ALIASES)

    print(f"     列検出: context={ci_context} hanzi={ci_hanzi} pinyin={ci_pinyin} "
          f"pos={ci_pos} ja={ci_ja} en={ci_en}")

    if ci_hanzi is None:
        print(f"  ⚠️  漢字列が見つかりません（シート: {sheet.title}）。スキップします。")
        print(f"     ヘッダー: {header}")
        return []

    def cell(row, idx) -> str:
        if idx is None or idx >= len(row) or row[idx] is None:
            return ""
        return str(row[idx]).strip()

    words = []
    word_id = start_id
    for row in rows[header_idx + 1:]:
        if not row:
            continue
        hanzi = strip_bopomofo(cell(row, ci_hanzi))
        if not hanzi:
            continue

        context = cell(row, ci_context)
        category = context_to_category(context)
        pinyin = cell(row, ci_pinyin)
        pos_raw = cell(row, ci_pos)

        words.append({
            "id":            word_id,
            "hanzi":         hanzi,
            "pinyin":        pinyin,
            "pinyin_tones":  to_pinyin_tones(pinyin) if pinyin else "",
            "meaning_ja":    cell(row, ci_ja),
            "meaning_en":    cell(row, ci_en),
            "part_of_speech": normalize_pos(pos_raw) if pos_raw else "noun",
            "category":      category,
            "context":       context,         # 元の任務領域を保持
            "frequency_rank": word_id,
            "difficulty":    difficulty,
            "tocfl_level":   tocfl_level,
            "example_sentence": {"hanzi": "", "pinyin": "", "meaning_ja": ""},
            "notes_ja":      "",
            "taiwan_specific": False,
        })
        word_id += 1

    return words

# ────────────────────────────────────────────────
# メイン
# ────────────────────────────────────────────────

def main() -> None:
    parser = argparse.ArgumentParser(description="Excel → vocab JSON 変換")
    parser.add_argument("--xlsx",  default="華語八千詞表_master.xlsx", help="Excelファイルのパス")
    parser.add_argument("--sheet", default=None,                      help="対象シート名（省略時は全シート）")
    parser.add_argument("--out",   default=None,                      help="出力JSONパス（省略時は自動分割）")
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx)
    if not xlsx_path.exists():
        sys.exit(f"❌  Excelファイルが見つかりません: {xlsx_path}")

    print(f"📂  {xlsx_path} を読み込み中...")
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    print(f"   シート一覧: {wb.sheetnames}")

    words: list[dict] = []
    word_id = 1

    for sheet in wb.worksheets:
        if args.sheet and sheet.title != args.sheet:
            continue
        diff = sheet_difficulty(sheet.title)
        tocfl = sheet_tocfl_level(sheet.title)
        print(f"   処理中: 「{sheet.title}」 (difficulty={diff}, tocfl_level={tocfl})")
        sheet_words = parse_sheet(sheet, diff, word_id, tocfl)
        print(f"   → {len(sheet_words)} 語")
        words.extend(sheet_words)
        word_id += len(sheet_words)

    wb.close()

    if not words:
        sys.exit("❌  単語が1件も取得できませんでした。")

    print(f"\n合計 {len(words)} 語を読み込みました")

    out_dir = Path("data")

    if args.out:
        out_path = Path(args.out)
        out_path.write_text(json.dumps(words, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"✅  {out_path}: {len(words)} 語")
    else:
        # 自動分割: 1-300, 301-600, 601+
        chunks = [
            ("vocab-core.json",     words[:300]),
            ("vocab-everyday.json", words[300:600]),
            ("vocab-advanced.json", words[600:]),
        ]
        for fname, chunk in chunks:
            if chunk:
                path = out_dir / fname
                path.write_text(json.dumps(chunk, ensure_ascii=False, indent=2), encoding="utf-8")
                print(f"✅  {path}: {len(chunk)} 語")


if __name__ == "__main__":
    main()
